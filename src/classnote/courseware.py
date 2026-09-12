from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path


MAX_SOURCE_CHARS = 80_000
MAX_HOTWORDS = 120
STOPWORDS = {
    "about", "after", "also", "and", "are", "because", "before", "between",
    "but", "can", "class", "course", "does", "during", "each", "for", "from",
    "have", "into", "its", "lecture", "more", "not", "our", "should", "that",
    "the", "their", "there", "these", "they", "this", "through", "today", "using",
    "was", "were", "what", "when", "where", "which", "will", "with", "would", "you",
}


@dataclass
class CourseContext:
    source_count: int = 0
    extracted_chars: int = 0
    hotwords: tuple[str, ...] = ()
    terms: dict[str, str] = field(default_factory=dict)
    warnings: tuple[str, ...] = ()

    @property
    def hotword_prompt(self) -> str:
        return ", ".join(self.hotwords)


def merge_subject_context(courseware: CourseContext, remembered: dict[str, str]) -> CourseContext:
    """Add persistent subject terms without changing the imported courseware context."""
    terms: dict[str, str] = {}
    seen_terms: set[str] = set()
    for source in (courseware.terms, remembered):
        for english, chinese in source.items():
            key = " ".join(english.split()).casefold()
            if not key or key in seen_terms or not chinese.strip():
                continue
            terms[english] = chinese
            seen_terms.add(key)
            if len(terms) >= 80:
                break
        if len(terms) >= 80:
            break
    hotwords: list[str] = []
    seen_hotwords: set[str] = set()
    for value in (*terms.keys(), *courseware.hotwords):
        key = " ".join(value.split()).casefold()
        if not key or key in seen_hotwords:
            continue
        hotwords.append(value)
        seen_hotwords.add(key)
        if len(hotwords) >= MAX_HOTWORDS:
            break
    return CourseContext(
        source_count=courseware.source_count,
        extracted_chars=courseware.extracted_chars,
        hotwords=tuple(hotwords),
        terms=terms,
        warnings=courseware.warnings,
    )


def relevant_terms(text: str, terms: dict[str, str], limit: int = 12) -> dict[str, str]:
    """Send only glossary entries present in this speech segment to the translator."""
    matches: list[tuple[int, int, str, str]] = []
    for english, chinese in terms.items():
        pattern = rf"(?<![A-Za-z0-9]){re.escape(english)}(?![A-Za-z0-9])"
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            matches.append((match.start(), -len(english), english, chinese))
    matches.sort()
    return {english: chinese for _, _, english, chinese in matches[:max(0, limit)]}


def build_course_context(paths: list[str]) -> CourseContext:
    texts: list[str] = []
    warnings: list[str] = []
    source_names: list[str] = []
    for raw in paths:
        path = Path(raw)
        source_names.append(path.stem)
        try:
            text = extract_courseware_text(path)
        except Exception as exc:
            warnings.append(f"{path.name}：{exc}")
            continue
        if text.strip():
            remaining = MAX_SOURCE_CHARS - sum(len(value) for value in texts)
            if remaining <= 0:
                warnings.append("课件内容较多，只提取了前 80000 个字符。")
                break
            texts.append(text[:remaining])
        else:
            warnings.append(f"{path.name}：未提取到可读文字，可能是扫描版课件")
    combined = "\n".join(source_names + texts)
    terms = extract_glossary(combined)
    hotwords = extract_hotwords(combined, terms)
    return CourseContext(
        source_count=len(paths),
        extracted_chars=sum(len(value) for value in texts),
        hotwords=tuple(hotwords),
        terms=terms,
        warnings=tuple(dict.fromkeys(warnings)),
    )


def extract_courseware_text(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError("找不到文件")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _read_text(path)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts: list[str] = []
        size = 0
        for index, page in enumerate(reader.pages):
            if index >= 160:
                break
            value = page.extract_text() or ""
            parts.append(value)
            size += len(value)
            if size >= MAX_SOURCE_CHARS:
                break
        return "\n".join(parts)[:MAX_SOURCE_CHARS]
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(str(path))
        parts: list[str] = []
        size = 0
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = str(shape.text)
                    parts.append(value)
                    size += len(value)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        value = " | ".join(cell.text for cell in row.cells)
                        parts.append(value)
                        size += len(value)
                if size >= MAX_SOURCE_CHARS:
                    break
            if size >= MAX_SOURCE_CHARS:
                break
        return "\n".join(parts)[:MAX_SOURCE_CHARS]
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        parts: list[str] = []
        size = 0
        for paragraph in document.paragraphs:
            parts.append(paragraph.text)
            size += len(paragraph.text)
            if size >= MAX_SOURCE_CHARS:
                break
        for table in document.tables:
            if size >= MAX_SOURCE_CHARS:
                break
            for row in table.rows:
                value = " | ".join(cell.text for cell in row.cells)
                parts.append(value)
                size += len(value)
                if size >= MAX_SOURCE_CHARS:
                    break
        return "\n".join(parts)[:MAX_SOURCE_CHARS]
    if suffix in {".ppt", ".doc"}:
        raise ValueError("旧版 PPT/DOC 暂不支持解析，请另存为 PPTX/DOCX")
    raise ValueError(f"暂不支持解析 {suffix or '无扩展名'} 文件")


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_glossary(text: str) -> dict[str, str]:
    terms: dict[str, str] = {}
    pattern = re.compile(
        r"^\s*([A-Za-z][A-Za-z0-9 .+#_()/\-]{1,64}?)\s*(?:[:：=]|\s[-–—]\s)\s*"
        r"([\u3400-\u9fff][^\r\n|]{0,48})\s*$"
    )
    for line in text.splitlines():
        match = pattern.match(line)
        if not match:
            continue
        english = " ".join(match.group(1).split()).strip(" -")
        chinese = match.group(2).strip(" 。；;")
        if 2 <= len(english) <= 64 and chinese:
            terms.setdefault(english, chinese)
        if len(terms) >= 80:
            break
    return terms


def extract_hotwords(text: str, terms: dict[str, str] | None = None) -> list[str]:
    terms = terms or {}
    tokens = re.findall(r"[A-Za-z][A-Za-z0-9.+#/_-]{2,}", text)
    counts = Counter(token.casefold() for token in tokens)
    representative: dict[str, str] = {}
    for token in tokens:
        key = token.casefold()
        current = representative.get(key)
        if current is None or (token.isupper() and not current.isupper()):
            representative[key] = token

    ranked: list[str] = list(terms.keys())
    # Short headings often contain the exact multi-word concept used by a lecturer.
    for line in text.splitlines():
        cleaned = re.sub(r"^[\s•·▪◦*#>\-–—\d.)]+", "", line).strip()
        words = re.findall(r"[A-Za-z][A-Za-z0-9.+#/_-]{1,}", cleaned)
        if (
            2 <= len(words) <= 7
            and len(cleaned) <= 90
            and not cleaned.endswith((".", "?", "!", ";"))
            and words[0].casefold() not in STOPWORDS
            and words[-1].casefold() not in STOPWORDS
        ):
            ranked.append(" ".join(words))
    acronyms = sorted(
        {token for token in tokens if token.isupper() and 2 <= len(token) <= 16},
        key=lambda value: (-counts[value.casefold()], value),
    )
    ranked.extend(acronyms)

    scored = sorted(
        (
            (count + (2 if any(char in token for char in ".+#/_-") else 0), token)
            for key, count in counts.items()
            if key not in STOPWORDS and len(key) >= 4
            for token in [representative[key]]
        ),
        key=lambda value: (-value[0], value[1].casefold()),
    )
    ranked.extend(token for _, token in scored)

    output: list[str] = []
    seen: set[str] = set()
    for value in ranked:
        normalized = " ".join(value.split()).strip(" ,.;:")
        key = normalized.casefold()
        if not normalized or key in seen:
            continue
        seen.add(key)
        output.append(normalized)
        if len(output) >= MAX_HOTWORDS:
            break
    return output
