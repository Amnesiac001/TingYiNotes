from pathlib import Path
from types import SimpleNamespace

import numpy as np
from docx import Document
from pptx import Presentation

from classnote.courseware import (
    CourseContext,
    build_course_context,
    extract_courseware_text,
    extract_glossary,
    extract_hotwords,
)
from classnote.local_live import LocalLiveCourseSession


def test_text_courseware_builds_glossary_and_multiword_hotwords(tmp_path: Path) -> None:
    source = tmp_path / "TCP congestion control.md"
    source.write_text(
        """# TCP Congestion Control
TCP：传输控制协议
Congestion Window - 拥塞窗口
Slow Start：慢启动

The congestion window grows during slow start.
""",
        encoding="utf-8",
    )

    context = build_course_context([str(source)])
    assert context.source_count == 1
    assert context.terms["TCP"] == "传输控制协议"
    assert context.terms["Congestion Window"] == "拥塞窗口"
    assert "tcp congestion control" in {value.casefold() for value in context.hotwords}
    assert "TCP" in context.hotwords
    assert not context.warnings


def test_docx_and_pptx_text_are_extracted(tmp_path: Path) -> None:
    docx_path = tmp_path / "terms.docx"
    document = Document()
    document.add_paragraph("Round Trip Time：往返时间")
    document.save(docx_path)

    pptx_path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Additive Increase Multiplicative Decrease"
    slide.placeholders[1].text = "AIMD：加性增乘性减"
    presentation.save(pptx_path)

    assert "Round Trip Time" in extract_courseware_text(docx_path)
    assert "Additive Increase" in extract_courseware_text(pptx_path)
    context = build_course_context([str(docx_path), str(pptx_path)])
    assert context.source_count == 2
    assert context.terms["Round Trip Time"] == "往返时间"
    assert "AIMD" in context.hotwords


def test_unreadable_courseware_warns_without_blocking_context() -> None:
    context = build_course_context(["D:/missing/lecture.pdf"])
    assert context.source_count == 1
    assert context.warnings


def test_empty_or_scanned_courseware_explains_missing_text(tmp_path: Path) -> None:
    empty = tmp_path / "scanned-notes.txt"
    empty.write_text("", encoding="utf-8")
    context = build_course_context([str(empty)])
    assert any("未提取到可读文字" in warning for warning in context.warnings)


def test_hotwords_are_forwarded_to_local_whisper() -> None:
    class FakeModel:
        kwargs: dict[str, object]

        def transcribe(self, audio: np.ndarray, **kwargs):
            self.kwargs = kwargs
            return [SimpleNamespace(text=" TCP congestion control ")], None

    session = object.__new__(LocalLiveCourseSession)
    session.course_context = CourseContext(hotwords=("TCP", "congestion window"))
    model = FakeModel()
    text = session._transcribe(model, np.zeros(16000, dtype=np.float32))

    assert text == "TCP congestion control"
    assert model.kwargs["hotwords"] == "TCP, congestion window"


def test_glossary_and_hotword_helpers_ignore_unhelpful_common_words() -> None:
    terms = extract_glossary("RTT: 往返时间\nnot a glossary sentence.")
    hotwords = extract_hotwords("the the the TCP latency latency", terms)
    assert terms == {"RTT": "往返时间"}
    assert "the" not in {value.casefold() for value in hotwords}
